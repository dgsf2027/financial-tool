import importlib.util
import pathlib
import unittest
from openpyxl import Workbook

SPEC = importlib.util.spec_from_file_location('t4_suite_dates', pathlib.Path(__file__).parents[1] / 'suite' / 'build_suite.py')
suite = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(suite)

class SuiteDateRangeTest(unittest.TestCase):
    def test_channel_headers_use_selected_dates_and_keep_range_values(self):
        wb = Workbook()
        meta = {'period': '2026-09', 'days': 2, 'dates': ['2026-09-12', '2026-09-13'],
                'rangeLabel': '2026-09-12 ～ 2026-09-13', 'inputKeys': ['retailIncome'],
                'metrics': [{'k': 'retailIncome', 'n': '零售收入', 'lvl': 1}]}
        ch = {'sheet': '渠道', 'name': '渠道', 'project': '橘农项目', 'buName': '橘农事业部', 'filled': 2}
        ws = suite.write_channel(wb, ch, [{'has': True, 'retailIncome': 20}, {'has': True, 'retailIncome': 0}], meta, {'retailIncome': 6}, 6)
        self.assertEqual(ws['C5'].value, '2026-09-12')
        self.assertEqual(ws['D5'].value, '2026-09-13')
        self.assertEqual(ws['C6'].value, 20)
        self.assertEqual(ws['D6'].value, 0)
        self.assertEqual(ws['B6'].value, '=SUM(C6:D6)')
        self.assertIn('2026-09-12 ～ 2026-09-13', ws['A1'].value)
        self.assertEqual(ws.max_column, 4)

    def test_legacy_month_payload_still_uses_month_day_headers(self):
        meta = {'period': '2026-09', 'days': 1, 'inputKeys': ['retailIncome'],
                'metrics': [{'k': 'retailIncome', 'n': '零售收入', 'lvl': 1}]}
        ch = {'sheet': '渠道', 'name': '渠道', 'project': '橘农项目', 'buName': '橘农事业部', 'filled': 0}
        ws = suite.write_channel(Workbook(), ch, [{'has': False}], meta, {'retailIncome': 6}, 6)
        self.assertEqual(ws['C5'].value, '1日')
        self.assertIsNone(ws['C6'].value)

if __name__ == '__main__':
    unittest.main()
