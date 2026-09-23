import contextlib
import importlib.util
import io
import json
import pathlib
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('suite_custom_expense', ROOT / 'suite' / 'build_suite.py')
suite = importlib.util.module_from_spec(spec)
spec.loader.exec_module(suite)


class CustomExpenseSuiteTest(unittest.TestCase):
    def payload(self):
        inputs = ['retailIncome', 'returnAmount', 'refundAmount', 'retailCost', 'returnCost'] + suite.OPERATING + ['expense_live'] + suite.DIRECT + suite.INDIRECT
        computed = ['salesIncome', 'salesCost', 'grossProfit', 'grossMargin', 'operating', 'direct', 'contribution', 'contributionRate', 'indirect', 'netProfit', 'netMargin']
        metrics = [{'k': key, 'n': '直播服务费' if key == 'expense_live' else key, 'lvl': 1 if key in inputs else 0, 'pct': key in suite.PCT_KEYS} for key in inputs + computed]
        day = {'has': True, 'retailIncome': 0, 'expense_live': 125.5}
        month = suite.derive(day, suite.OPERATING + ['expense_live'])
        return {'period': '2026-09', 'days': 1, 'scopeName': '全部项目', 'generated': '测试',
                'metrics': metrics, 'inputKeys': inputs, 'operatingKeys': suite.OPERATING + ['expense_live'],
                'channels': [{'id': 'shop', 'name': '测试店', 'buName': '测试部', 'project': '测试项目', 'filled': 0}],
                'tree': [{'name': '全部', 'children': [{'name': '测试部', 'children': [{'name': '测试店', 'id': 'shop', 'children': []}]}]}],
                'dailyByCh': {'shop': [day]}, 'monthByCh': {'shop': month}}

    def test_custom_expense_is_in_formulas_at_every_report_level(self):
        payload = self.payload()
        with tempfile.TemporaryDirectory() as directory:
            inp, out = pathlib.Path(directory) / 'in.json', pathlib.Path(directory) / 'out.xlsx'
            inp.write_text(json.dumps(payload), encoding='utf-8')
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                suite.main(str(inp), str(out))
            self.assertTrue(json.loads(output.getvalue())['ok'])
            workbook = load_workbook(out, data_only=False)
            rows = {m['k']: 6 + index for index, m in enumerate(payload['metrics'])}
            custom_row, operating_row = rows['expense_live'], rows['operating']
            self.assertEqual(workbook['测试店'].cell(custom_row, 3).value, 125.5)
            for sheet in ['总表', '测试部', '测试店']:
                self.assertIn(f'B{custom_row}', workbook[sheet].cell(operating_row, 2).value)
            self.assertIn(f'C{custom_row}', workbook['测试店'].cell(operating_row, 3).value)
            self.assertEqual(workbook['测试店'].cell(rows['contribution'], 3).value,
                             f'=C{rows["grossProfit"]}-C{operating_row}-C{rows["direct"]}')

    def test_legacy_payload_uses_fixed_operating_expenses(self):
        payload = self.payload()
        del payload['operatingKeys']
        self.assertEqual(suite.operating_keys(payload), suite.OPERATING)
        self.assertEqual(suite.derive({'platformFee': 15})['netProfit'], -15)

    def test_malformed_dynamic_key_cannot_become_an_operating_formula(self):
        payload = self.payload()
        for invalid in ['netProfit', 'missing', 'expense_bad[]', 42]:
            payload['operatingKeys'] = suite.OPERATING + [invalid]
            with self.assertRaises(ValueError):
                suite.operating_keys(payload)


if __name__ == '__main__':
    unittest.main()
