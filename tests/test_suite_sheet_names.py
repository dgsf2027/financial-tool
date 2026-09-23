"""Workbook references remain distinct for Excel's case-insensitive tab names."""
import contextlib
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

SPEC = importlib.util.spec_from_file_location('t4_suite_names', Path(__file__).parents[1] / 'suite' / 'build_suite.py')
suite = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(suite)


class SuiteSheetNamesTest(unittest.TestCase):
    def test_colliding_channel_and_department_names_keep_their_own_references(self):
        channels = [
            {'id': 'one', 'name': 'Store', 'project': '项目', 'buName': 'STORE', 'filled': 1},
            {'id': 'two', 'name': 'store', 'project': '项目', 'buName': 'STORE', 'filled': 1},
        ]
        data = {
            'period': '2026-09', 'days': 1, 'scopeName': '测试', 'generated': '测试',
            'metrics': [{'k': 'retailIncome', 'n': '零售收入', 'lvl': 1}],
            'inputKeys': ['retailIncome'], 'channels': channels,
            'tree': [{'name': '项目', 'children': [{'name': 'STORE', 'children': [
                {'name': c['name'], 'id': c['id'], 'children': []} for c in channels]}]}],
            'dailyByCh': {'one': [{'has': True, 'retailIncome': 100}],
                          'two': [{'has': True, 'retailIncome': 200}]},
            'monthByCh': {key: dict(salesIncome=value, grossProfit=value, netProfit=value, operating=0)
                          for key, value in [('one', 100), ('two', 200)]},
        }
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'input.json', Path(directory) / 'suite.xlsx'
            source.write_text(json.dumps(data), encoding='utf-8')
            with contextlib.redirect_stdout(io.StringIO()):
                suite.main(str(source), str(output))
            wb = load_workbook(output)
            try:
                # Identify each actual tab through its title; openpyxl silently
                # renames duplicate titles, so a uniqueness assertion alone misses the bug.
                store = next(ws for ws in wb if ws['A1'].value.startswith('Store ·'))
                other = next(ws for ws in wb if ws['A1'].value.startswith('store ·'))
                department = next(ws for ws in wb if ws['A1'].value.startswith('STORE ·'))
                self.assertEqual(store['C6'].value, 100)
                self.assertEqual(other['C6'].value, 200)
                for column, channel in [('C', store), ('D', other)]:
                    self.assertEqual(department[f'{column}6'].value, f"='{channel.title}'!B6")
                    self.assertEqual(department[f'{column}6'].hyperlink.target, f"#'{channel.title}'!B6")
                    self.assertEqual(channel['B3'].hyperlink.target, f"#'{department.title}'!A1")
                self.assertEqual(wb['总表']['C6'].value, f"='{store.title}'!B6+'{other.title}'!B6")
                self.assertEqual(wb['总表']['C6'].hyperlink.target, f"#'{department.title}'!B6")
            finally:
                wb.close()

    def test_case_variants_after_invalid_character_replacement_remain_unique(self):
        used = {'总表'}
        names = [suite.sheet_name(name, used) for name in ['Shop/Online', 'shop\\online', 'SHOP·ONLINE~2']]
        self.assertEqual(len({name.lower() for name in names}), len(names))


if __name__ == '__main__':
    unittest.main()
