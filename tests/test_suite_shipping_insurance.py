import contextlib
import importlib.util
import io
import json
import pathlib
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('suite_shipping_insurance', ROOT / 'suite' / 'build_suite.py')
suite = importlib.util.module_from_spec(spec)
spec.loader.exec_module(suite)


class ShippingInsuranceSuiteTest(unittest.TestCase):
    def payload(self, insurance=True):
        operating = [key for key in suite.OPERATING if insurance or key != 'shippingInsurance']
        inputs = ['retailIncome', 'returnAmount', 'refundAmount', 'retailCost', 'returnCost'] + operating + suite.DIRECT + suite.INDIRECT
        computed = ['salesIncome', 'salesCost', 'grossProfit', 'grossMargin', 'operating', 'direct', 'contribution', 'contributionRate', 'indirect', 'netProfit', 'netMargin']
        metrics = [{'k': key, 'n': '运费险' if key == 'shippingInsurance' else key, 'lvl': 1 if key in inputs else 0, 'pct': key in suite.PCT_KEYS} for key in inputs + computed]
        day = {'has': True, 'retailIncome': 1000, 'retailCost': 400, 'platformFee': 50}
        if insurance:
            day['shippingInsurance'] = 12.5
        return {'period': '2026-09', 'days': 1, 'scopeName': '全部项目', 'generated': '测试',
                'metrics': metrics, 'inputKeys': inputs, 'operatingKeys': operating,
                'channels': [{'id': 'dycreator', 'name': '抖音-BD达人成交店', 'buName': '经销事业部', 'project': '澳乐项目', 'filled': 1}],
                'tree': [{'name': '全部', 'children': [{'name': '经销事业部', 'children': [{'name': '抖音-BD达人成交店', 'id': 'dycreator', 'children': []}]}]}],
                'dailyByCh': {'dycreator': [day]}, 'monthByCh': {'dycreator': suite.derive(day, operating)}}

    def build(self, payload, directory):
        inp, out = pathlib.Path(directory) / 'in.json', pathlib.Path(directory) / 'out.xlsx'
        inp.write_text(json.dumps(payload), encoding='utf-8')
        with contextlib.redirect_stdout(io.StringIO()):
            suite.main(str(inp), str(out))
        return load_workbook(out, data_only=False)

    def test_insurance_is_in_operating_formulas_at_every_report_level(self):
        payload = self.payload()
        self.assertEqual(payload['monthByCh']['dycreator']['netProfit'], 537.5)
        with tempfile.TemporaryDirectory() as directory:
            workbook = self.build(payload, directory)
            rows = {metric['k']: 6 + index for index, metric in enumerate(payload['metrics'])}
            insurance_row, operating_row = rows['shippingInsurance'], rows['operating']
            self.assertEqual(workbook['抖音-BD达人成交店'].cell(insurance_row, 3).value, 12.5)
            for name in ['总表', '经销事业部', '抖音-BD达人成交店']:
                self.assertIn(f'B{insurance_row}', workbook[name].cell(operating_row, 2).value)
            self.assertIn(f'C{insurance_row}', workbook['抖音-BD达人成交店'].cell(operating_row, 3).value)
            self.assertEqual(workbook['抖音-BD达人成交店'].cell(rows['contribution'], 3).value,
                             f'=C{rows["grossProfit"]}-C{operating_row}-C{rows["direct"]}')
            workbook.close()

    def test_existing_payload_without_insurance_row_still_builds(self):
        payload = self.payload(insurance=False)
        del payload['operatingKeys']
        self.assertNotIn('shippingInsurance', suite.operating_keys(payload))
        with tempfile.TemporaryDirectory() as directory:
            workbook = self.build(payload, directory)
            self.assertEqual(len(workbook.worksheets), 3)
            self.assertFalse(any(cell.value == '运费险' for sheet in workbook for row in sheet for cell in row))
            workbook.close()

    def test_python_profit_recalculation_includes_insurance_actuals(self):
        self.assertEqual(suite.derive({'shippingInsurance': 25})['netProfit'], -25)
        self.assertEqual(suite.derive({'shippingInsurance': 0})['netProfit'], 0)
        self.assertEqual(suite.derive({'shippingInsurance': -5})['netProfit'], 5)


if __name__ == '__main__':
    unittest.main()
