import contextlib
import importlib.util
import io
import json
import pathlib
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('suite_returns', ROOT / 'suite' / 'build_suite.py')
suite = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(suite)


class ReturnsSuiteTest(unittest.TestCase):
    def payload(self):
        inputs = (['retailIncome', 'returnAmount', 'refundAmount', 'rebateAmount', 'retailCost', 'returnCost',
                   'rebateIncome', 'salesReceipt'] + suite.OPERATING + ['expense_live']
                  + suite.DIRECT + suite.INDIRECT)
        calculated = ['salesIncome', 'salesCost', 'grossProfit', 'grossMargin', 'operating',
                      'direct', 'contribution', 'contributionRate', 'indirect', 'netProfit', 'netMargin']
        metrics = [{'k': key, 'n': key, 'lvl': 1 if key in inputs else 0,
                    'pct': key in suite.PCT_KEYS} for key in inputs + calculated]
        days = [
            {'has': True, 'retailIncome': 1000, 'refundAmount': -100, 'rebateAmount': 100, 'retailCost': 400,
             'platformFee': 40, 'expense_live': 25, 'directLabor': 10, 'sharedLabor': 5,
             'rebateIncome': 30, 'salesReceipt': 800},
            {'has': True, 'rebateAmount': -100, 'rebateIncome': 20, 'salesReceipt': 200,
             'platformFee': -5, 'expense_live': 10},
            {'has': True, 'rebateAmount': 0, 'rebateIncome': 0, 'salesReceipt': 0, 'expense_live': 0},
        ]
        return {'period': '2026-09', 'days': 3,
                'dates': ['2026-09-12', '2026-09-13', '2026-09-14'],
                'rangeLabel': '2026-09-12 ～ 2026-09-14', 'scopeName': '橘农项目', 'generated': '测试',
                'metrics': metrics, 'inputKeys': inputs,
                'operatingKeys': suite.OPERATING + ['expense_live'],
                'channels': [{'id': 'shop', 'name': '返款测试店', 'buName': '橘农事业部',
                              'project': '橘农项目', 'filled': 1}],
                'tree': [{'name': '橘农项目', 'children': [
                    {'name': '橘农事业部', 'children': [
                        {'name': '返款测试店', 'id': 'shop', 'children': []}]}]}],
                'dailyByCh': {'shop': days},
                # Both signs deduct independently: 1000 - 100 - 100 - 100 = 700.
                # Frontend's 5% platform fees are inputs: 40 - 5 = 35.
                # Net profit: 700 - 400 - (35 + 35) - 10 - 5 + 50 = 265.
                'monthByCh': {'shop': {'salesIncome': 700, 'grossProfit': 300,
                                      'operating': 70, 'netProfit': 265}}}

    def test_xlsx_return_formulas_at_channel_business_and_project_levels(self):
        payload = self.payload()
        with tempfile.TemporaryDirectory() as directory:
            source, output = pathlib.Path(directory) / 'input.json', pathlib.Path(directory) / 'returns.xlsx'
            source.write_text(json.dumps(payload), encoding='utf-8')
            report = io.StringIO()
            with contextlib.redirect_stdout(report):
                suite.main(str(source), str(output))
            self.assertEqual(json.loads(report.getvalue())['verify_mismatch'], [])
            self.assertTrue(json.loads(report.getvalue())['ok'])
            wb = load_workbook(output, data_only=False)
            self.addCleanup(wb.close)
            rows = {metric['k']: 6 + index for index, metric in enumerate(payload['metrics'])}

            for sheet_name, columns in [('总表', ['B', 'C']), ('橘农事业部', ['B']),
                                         ('返款测试店', ['B', 'C', 'D', 'E'])]:
                for col in columns:
                    with self.subTest(sheet=sheet_name, col=col):
                        cell = lambda key: f'{col}{rows[key]}'
                        ws = wb[sheet_name]
                        # New returns deduct sales separately from refunds; old income/cash stay independent.
                        self.assertEqual(ws[cell('netProfit')].value,
                                         f'={cell("contribution")}-{cell("indirect")}+{cell("rebateIncome")}')
                        self.assertEqual(ws[cell('salesIncome')].value,
                                         f'={cell("retailIncome")}+{cell("returnAmount")}+{cell("refundAmount")}-ABS({cell("rebateAmount")})')
                        self.assertIn(cell('expense_live'), ws[cell('operating')].value)
                        for key in ['salesIncome', 'grossProfit', 'operating', 'contribution', 'netProfit']:
                            self.assertNotIn(cell('salesReceipt'), ws[cell(key)].value)

            channel = wb['返款测试店']
            for key, expected in [('refundAmount', [-100, None, None]),
                                  ('rebateAmount', [-100, -100, 0]),
                                  ('rebateIncome', [30, 20, 0]),
                                  ('salesReceipt', [800, 200, 0]),
                                  ('expense_live', [25, 10, 0])]:
                self.assertEqual([channel.cell(rows[key], col).value for col in (3, 4, 5)], expected)
                self.assertEqual(channel.cell(rows[key], 2).value,
                                 f'=-SUMPRODUCT(ABS(C{rows[key]}:E{rows[key]}))' if key == 'rebateAmount'
                                 else f'=SUM(C{rows[key]}:E{rows[key]})')
            # Suite uses the frontend's fee values without recalculating from retail income.
            self.assertEqual([channel.cell(rows['platformFee'], col).value for col in (3, 4)], [40, -5])
            self.assertEqual(wb['橘农事业部'].cell(rows['salesReceipt'], 3).value,
                             f"='返款测试店'!B{rows['salesReceipt']}")
            self.assertEqual(wb['总表'].cell(rows['rebateIncome'], 3).value,
                             f"='返款测试店'!B{rows['rebateIncome']}")
            self.assertEqual(wb['总表'].cell(rows['rebateAmount'], 3).value,
                             f"='返款测试店'!B{rows['rebateAmount']}")
            report_note = wb['总表'].cell(max(rows.values()) + 3, 1).value
            self.assertIn('退款金额-返款金额绝对值', report_note)
            self.assertIn('费用比例统一按销售收入计算，实填金额优先', report_note)
            self.assertTrue(wb.calculation.fullCalcOnLoad)

    def test_python_verification_keeps_cash_out_of_profit_with_custom_expenses(self):
        payload = self.payload()
        operating = payload['operatingKeys']
        for day, expected in zip(payload['dailyByCh']['shop'], [350, -85, 0]):
            with self.subTest(day=day):
                result = suite.derive(day, operating)
                self.assertEqual(result['netProfit'], expected)
                self.assertEqual(suite.derive({**day, 'salesReceipt': 999999}, operating)['netProfit'], expected)
        zero_rebate = suite.derive({'rebateIncome': 0, 'salesReceipt': 800, 'expense_live': 12}, operating)
        self.assertEqual(zero_rebate['salesIncome'], 0)
        self.assertEqual(zero_rebate['netProfit'], -12)

    def test_positive_and_negative_returns_deduct_without_overwriting_refunds(self):
        payload = self.payload()
        day = payload['dailyByCh']['shop'][0]
        for value in [100, -100]:
            with self.subTest(rebateAmount=value):
                result = suite.derive({**day, 'rebateAmount': value}, payload['operatingKeys'])
                self.assertEqual(result['rebateAmount'], -100)
                self.assertEqual(result['refundAmount'], -100)
                self.assertEqual(result['salesIncome'], 800)
                self.assertEqual(result['platformFee'], 40)
                self.assertEqual(result['netProfit'], 350)
                self.assertEqual(result['rebateIncome'], 30)
                self.assertEqual(result['salesReceipt'], 800)

    def test_legacy_payload_without_return_field_keeps_compatible_formulas(self):
        payload = self.payload()
        legacy_rows = {metric['k']: 6 + index for index, metric in enumerate(payload['metrics'])
                       if metric['k'] != 'rebateAmount'}
        r = lambda key: f'B{legacy_rows[key]}'
        self.assertEqual(suite.formula_for('salesIncome', 'B', legacy_rows),
                         f'={r("retailIncome")}+{r("returnAmount")}+{r("refundAmount")}')
        self.assertEqual(suite.formula_for('netProfit', 'B', legacy_rows),
                         f'={r("contribution")}-{r("indirect")}+{r("rebateIncome")}')
        del legacy_rows['rebateIncome']
        self.assertEqual(suite.formula_for('netProfit', 'B', legacy_rows),
                         f'={r("contribution")}-{r("indirect")}')
        legacy = {'retailIncome': 1000, 'refundAmount': -100, 'retailCost': 400,
                  'platformFee': 45, 'expense_live': 25, 'rebateIncome': 30, 'salesReceipt': 800}
        calculated = suite.derive(legacy, payload['operatingKeys'])
        self.assertEqual(calculated['salesIncome'], 900)
        self.assertEqual(calculated['netProfit'], 460)


if __name__ == '__main__':
    unittest.main()
